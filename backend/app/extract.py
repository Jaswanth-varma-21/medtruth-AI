"""Turn an uploaded file into report text.

Supports: .txt/.md, .docx, text PDFs, scanned PDFs (OCR) and images (OCR).
OCR needs the Tesseract program (and Poppler for PDFs) installed on the machine; see README.
Note: OCR reads *printed words*. It does not interpret an X-ray picture itself.
"""
import io

MAX_BYTES = 5 * 1024 * 1024
MAX_CHARS = 20000
MAX_OCR_PAGES = 10
MIN_WORD_CONF = 60           # OCR words below this confidence (0-100) are discarded
MIN_WORDS = 3
MIN_TEXT_CHARS = 20          # fewer characters than this in a PDF text layer => treat as scanned
IMAGE_EXT = (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp")


class ExtractError(Exception):
    """A problem the user can understand and act on."""


def _clean(text: str) -> str:
    lines = [ln.rstrip() for ln in text.replace("\r", "").split("\n")]
    out = "\n".join(lines).strip()
    while "\n\n\n" in out:
        out = out.replace("\n\n\n", "\n\n")
    return out


def _ocr_image(img) -> str:
    try:
        import pytesseract
        from PIL import ImageOps
    except ImportError:
        raise ExtractError("OCR needs: pip install pytesseract pillow")
    img = ImageOps.exif_transpose(img).convert("L")
    img = ImageOps.autocontrast(img)
    if max(img.size) < 1500:                      # small images OCR badly; upscale
        img = img.resize((img.width * 2, img.height * 2))
    try:
        d = pytesseract.image_to_data(img, config="--psm 6", output_type=pytesseract.Output.DICT)
    except pytesseract.TesseractNotFoundError:
        raise ExtractError("Tesseract is not installed. In WSL run: sudo apt install -y tesseract-ocr")
    # Keep only confidently-read words; this drops the noise OCR invents from a plain X-ray picture.
    lines, order = {}, []
    for i, word in enumerate(d["text"]):
        try:
            conf = float(d["conf"][i])
        except ValueError:
            continue
        if word.strip() and conf >= MIN_WORD_CONF:
            key = (d["block_num"][i], d["par_num"][i], d["line_num"][i])
            if key not in lines:
                lines[key] = []
                order.append(key)
            lines[key].append(word)
    return "\n".join(" ".join(lines[k]) for k in order)


def _from_image(data: bytes) -> str:
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception:
        raise ExtractError("Could not open that image.")
    return _ocr_image(img)


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractError("That PDF is password protected.")
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
    except ExtractError:
        raise
    except ImportError:
        raise ExtractError("PDF reading needs: pip install pypdf")
    except Exception:
        raise ExtractError("Could not read that PDF.")
    if len(text.strip()) >= MIN_TEXT_CHARS:
        return text
    # No text layer: it is a scan. Render pages to images and OCR them.
    try:
        from pdf2image import convert_from_bytes
        from pdf2image.exceptions import PDFInfoNotInstalledError
    except ImportError:
        raise ExtractError("Scanned PDFs need: pip install pdf2image pytesseract pillow")
    try:
        pages = convert_from_bytes(data, dpi=200, first_page=1, last_page=MAX_OCR_PAGES)
    except PDFInfoNotInstalledError:
        raise ExtractError("Poppler is not installed. In WSL run: sudo apt install -y poppler-utils")
    except Exception:
        raise ExtractError("Could not render that PDF for OCR.")
    return "\n\n".join(_ocr_image(p) for p in pages)


def _from_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError:
        raise ExtractError("Word files need: pip install python-docx")
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception:
        raise ExtractError("Could not read that Word file (only .docx is supported, not old .doc).")
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append("  ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def extract_text(data: bytes, filename: str) -> str:
    if not data:
        raise ExtractError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise ExtractError("File is larger than 5 MB.")
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        text = _from_pdf(data)
    elif name.endswith(".docx"):
        text = _from_docx(data)
    elif name.endswith(IMAGE_EXT):
        text = _from_image(data)
    elif name.endswith((".txt", ".text", ".md")):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ExtractError("Unsupported file type. Use .pdf, .docx, .txt, .md, .png, .jpg, .tif, .bmp or .webp.")
    text = _clean(text)
    if len(text) < 10 or sum(1 for w in text.split() if sum(c.isalpha() for c in w) >= 3) < MIN_WORDS:
        raise ExtractError(
            "No readable text found. MedTruth reads the words of a report; it cannot interpret an X-ray picture itself. "
            "If this is a photo of a report, try a sharper, straighter picture."
        )
    return text[:MAX_CHARS]
